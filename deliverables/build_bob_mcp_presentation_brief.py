from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DOCX = Path(
    "/Users/bilolbekrayimov/games/IBM_hackathon/deliverables/bob-mcp-presentation-brief.docx"
)

TITLE_COLOR = RGBColor(11, 37, 69)  # #0B2545
H1_COLOR = RGBColor(46, 116, 181)  # #2E74B5
H2_COLOR = RGBColor(31, 77, 120)  # #1F4D78
BODY_COLOR = RGBColor(17, 24, 39)
MUTED_COLOR = RGBColor(98, 108, 123)


def set_run_font(run, *, name: str = "Calibri", size: int | None = None,
                 color: RGBColor | None = None, bold: bool | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_border_bottom(paragraph, color: str = "D7DEE8", size: str = "12") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)

    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)


def add_page_field(paragraph) -> None:
    fld_simple = OxmlElement("w:fldSimple")
    fld_simple.set(qn("w:instr"), "PAGE")

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld_simple.append(run)
    paragraph._p.append(fld_simple)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY_COLOR
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = TITLE_COLOR
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = "Calibri"
    heading1.font.size = Pt(16)
    heading1.font.bold = True
    heading1.font.color.rgb = H1_COLOR
    heading1.paragraph_format.space_before = Pt(16)
    heading1.paragraph_format.space_after = Pt(8)
    heading1.paragraph_format.line_spacing = 1.0

    heading2 = doc.styles["Heading 2"]
    heading2.font.name = "Calibri"
    heading2.font.size = Pt(13)
    heading2.font.bold = True
    heading2.font.color.rgb = H2_COLOR
    heading2.paragraph_format.space_before = Pt(12)
    heading2.paragraph_format.space_after = Pt(6)
    heading2.paragraph_format.line_spacing = 1.0


def build_header_footer(section) -> None:
    header = section.header
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_para.paragraph_format.space_after = Pt(0)
    header_run = header_para.add_run("Modernization Navigator | Presentation Brief")
    set_run_font(header_run, size=9, color=MUTED_COLOR)

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_para.paragraph_format.space_after = Pt(0)
    text_run = footer_para.add_run("Why This MCP Makes Bob Better | Page ")
    set_run_font(text_run, size=9, color=MUTED_COLOR)
    add_page_field(footer_para)


def add_meta_line(doc: Document, label: str, value: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.05
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=10, color=BODY_COLOR, bold=True)
    value_run = paragraph.add_run(value)
    set_run_font(value_run, size=10, color=BODY_COLOR)


def add_kicker(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text.upper())
    set_run_font(run, size=10, color=H1_COLOR, bold=True)


def add_subtitle(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(16)
    run = paragraph.add_run(text)
    set_run_font(run, size=12, color=MUTED_COLOR)


def add_emphasis_paragraph(doc: Document, label: str, body: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.12
    lead = paragraph.add_run(f"{label}. ")
    set_run_font(lead, size=11, color=BODY_COLOR, bold=True)
    body_run = paragraph.add_run(body)
    set_run_font(body_run, size=11, color=BODY_COLOR)


def add_quote_line(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(f"“{text}”")
    set_run_font(run, size=11, color=TITLE_COLOR, italic=True)


def build_document() -> Path:
    doc = Document()
    style_document(doc)
    build_header_footer(doc.sections[0])

    add_kicker(doc, "Design Support Brief")

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.keep_with_next = True
    title_run = title.add_run("Why This MCP Makes Bob Better")
    set_run_font(title_run, size=24, color=TITLE_COLOR, bold=True)

    add_subtitle(
        doc,
        "A concise messaging brief to help designers frame the product as Bob augmentation, not Bob replacement."
    )

    add_meta_line(doc, "Audience", "Presentation design team")
    add_meta_line(doc, "Objective", "Make the MCP feel like a force multiplier for Bob, not a system that takes Bob’s work away")
    add_meta_line(doc, "Source Basis", "Repository analysis, architecture docs, live tool registry, and passing build plus test verification")
    add_meta_line(doc, "Date", "May 16, 2026")

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(14)
    set_paragraph_border_bottom(rule)

    doc.add_paragraph(
        "Core message: This MCP does not do Bob’s job for him. It gives Bob repo-wide evidence, repeatable technical analysis, and durable outputs so Bob’s recommendations become faster, sharper, and more defensible."
    )

    doc.add_heading("Narrative to Land", level=1)
    doc.add_paragraph(
        "Bob remains the architect in the conversation. He still asks the questions, decides which tools to use, weighs the tradeoffs, and delivers the final recommendation. The MCP plays the supporting role: it scans the repository, extracts facts, scores paths, saves the artifact, and opens the viewer."
    )
    doc.add_paragraph(
        "That distinction matters for the story. The value is not autonomous migration. The value is stronger judgment. Bob stays human-centered and conversational, while the MCP makes the technical foundation under Bob much more reliable."
    )

    doc.add_heading("Strong Points Designers Can Build Around", level=1)

    add_emphasis_paragraph(
        doc,
        "1. Bob stays in charge",
        "The product contract is explicit that Bob owns the recommendation, the plan, and the user-facing explanation. The experience still feels like a strategist guiding the user, not a background script replacing the strategist."
    )
    add_emphasis_paragraph(
        doc,
        "2. Bob gains repo-wide x-ray vision",
        "The MCP scans package files, lockfiles, runtime pins, CI workflows, Dockerfiles, deployment configuration, and risky source patterns. Bob becomes more complete and less likely to miss hidden blockers outside the obvious package surface."
    )
    add_emphasis_paragraph(
        doc,
        "3. Bob gets repeatable technical evidence",
        "Instead of relying on vague AI intuition, the MCP returns typed findings, normalized issues, and validation steps. That makes Bob’s answers easier to trust, easier to review, and easier to reuse."
    )
    add_emphasis_paragraph(
        doc,
        "4. Bob can justify direct versus staged paths",
        "The MCP compares upgrade routes with explicit risk and effort scoring. Bob can explain why the safest path or the fastest path wins for this specific repository instead of sounding generic."
    )
    add_emphasis_paragraph(
        doc,
        "5. Bob creates durable artifacts, not disappearing chat",
        "Each run can become a saved report with decision, issues, evidence, execution plan, validation checklist, and tool trace. The output becomes something a team can revisit, share, and defend."
    )
    add_emphasis_paragraph(
        doc,
        "6. The system is safe by design",
        "This MCP is built to inspect and recommend without silently changing the analyzed repository. That makes the experience feel supportive and trustworthy instead of autonomous and opaque."
    )
    add_emphasis_paragraph(
        doc,
        "7. The capability scales beyond one repo",
        "The install flow lets the same MCP attach to other repositories, and the v2 architecture already opens a path beyond Node into React and future stacks. Bob becomes better everywhere this evidence layer is plugged in."
    )

    doc.add_heading("What Designers Should Emphasize", level=1)
    add_emphasis_paragraph(
        doc,
        "Best visual metaphor",
        "Use navigator, x-ray vision, instrument panel, or evidence engine. Those metaphors communicate augmentation. Avoid autonomous robot imagery, because it suggests the tool is replacing Bob rather than empowering Bob."
    )
    add_emphasis_paragraph(
        doc,
        "Best flow to show",
        "User prompt to Bob, Bob to MCP tools, MCP to evidence and scoring, then back to Bob for the recommendation, with a saved report and viewer as the reusable artifact. That flow keeps Bob at the center while showing why the MCP matters."
    )
    add_emphasis_paragraph(
        doc,
        "Best product language",
        "Lean on phrases like evidence-backed, defensible, repeatable, repo-aware, and decision support. Avoid language that makes the MCP sound like an autopilot or a code-writing replacement agent."
    )

    doc.add_heading("Slide-Ready Lines", level=1)
    add_quote_line(
        doc,
        "This MCP does not do Bob’s job for him. It makes Bob’s judgment faster, sharper, and more defensible."
    )
    add_quote_line(
        doc,
        "Bob stays the strategist; the MCP supplies the instrumentation."
    )
    add_quote_line(
        doc,
        "The value is not more automation. The value is better judgment backed by evidence."
    )

    doc.add_heading("Verification Note", level=1)
    doc.add_paragraph(
        "This brief is grounded in a full repository review. The project build passed and the automated test suite passed at the time of analysis, which supports the credibility of the messaging above."
    )

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    path = build_document()
    print(path)
